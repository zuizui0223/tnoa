#!/usr/bin/env python3
"""Source-only wind/pollinator file verification. No biological outcome contrasts.

Fetch original public Balfour and Coates data without copying originals to Git.
Report counts, column names, non-response source support, and SHA256 for later pinning.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
from pathlib import Path
import time
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from zipfile import ZipFile

SOURCES = {
    "balfour_zip": "https://datadryad.org/downloads/file_stream/4764144",
    "coates_file_a": "https://datadryad.org/downloads/file_stream/3519141",
    "coates_file_b": "https://datadryad.org/downloads/file_stream/3519140",
    "crall_zenodo_zip": "https://zenodo.org/records/3675327/files/CrallEtAl_OrchidBeeWindEcology.zip?download=1",
}
HEADERS = {
    "User-Agent": "Mozilla/5.0 (compatible; ecology-research-source-audit/1.0; +https://github.com/zuizui0223/tnoa)",
    "Accept": "*/*",
}

def download(url: str) -> tuple[bytes,str]:
    failures = []
    for attempt in range(1,4):
        try:
            with urlopen(Request(url, headers=HEADERS), timeout=45) as response:
                content = response.read()
                ctype = response.headers.get("Content-Type","")
                if len(content)<1000 or content[:7].lower().startswith(b"<!doct"):
                    raise ValueError(f"Unexpected tiny or HTML download, content_type={ctype}: {content[:150]!r}")
                return content,ctype
        except (HTTPError,URLError,ValueError,TimeoutError) as exc:
            failures.append(str(exc)[:230])
            if attempt<3: time.sleep(2)
    raise RuntimeError(f"Source failure for {url}: {failures}")

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--output",default="research/wind_process_overlap/intake_receipt.json")
    args=p.parse_args()
    receipt={
      "schema":"tnoa-wind-ecology-public-source-audit-v1",
      "date":"2026-10-19",
      "status":"source_only_no_biological_outcomes",
      "source_urls":SOURCES,
      "original_files":{},
      "source_errors":{},
      "original_publications":[
         {"dataset":"Balfour and Ratnieks","doi":"10.5061/dryad.4qrfj6qpf","article":"10.1093/beheco/araf067"},
         {"dataset":"Coates et al.","doi":"10.5061/dryad.0zpc86753","article":"10.1111/oik.10688"},
         {"dataset":"Crall et al.","doi":"10.5281/zenodo.3675327","article":"10.1098/rsbl.2020.0103"}],
      "boundaries":["no independent ecological effects opened","archive SHA256 not yet a permanent immutable upstream version","row counts are data structure not independent biological units","do not redistribute raw files"]
    }
    for key,url in SOURCES.items():
        try:
            content,ct = download(url)
            fileinfo={"source_url":url,"bytes":len(content),
                      "sha256":hashlib.sha256(content).hexdigest(),"http_content_type":ct}
            if not content.startswith(b"PK\x03\x04"):
                raise ValueError(f"{key}: expected ZIP/XLSX archive, received {content[:32]!r}")
            if key=="crall_zenodo_zip" and fileinfo["sha256"]!="d5b217261e4695e2098bde90576238018ad978c581f4b7799c15e429be19384b":
                raise ValueError("Crall original ZIP SHA256 drift; stop before opening data")
            with ZipFile(io.BytesIO(content)) as z:
                names=z.namelist()
                if key.endswith("_zip"):
                    fileinfo["archive_members"]=names
                    csvnames=[n for n in names if n.lower().endswith(".csv") and not n.startswith("__MACOSX")]
                    fileinfo["members"]={}
                    for name in csvnames:
                        raw=z.read(name)
                        text=raw.decode("utf-8-sig",errors="replace")
                        delimiter=csv.Sniffer().sniff(text[:4096],delimiters=",;\t").delimiter
                        table=list(csv.reader(io.StringIO(text),delimiter=delimiter))
                        fileinfo["members"][name.split("/")[-1]]={
                          "rows":len(table)-1,"columns":table[0],
                          "member_sha256":hashlib.sha256(raw).hexdigest()}
                    if key=="balfour_zip":
                        needed=["Surveys","Foraging","Inflorescence"]
                        if not all(any(t.lower() in zname.lower() for zname in fileinfo["members"]) for t in needed):
                            raise ValueError(f"Balfour unexpected CSV names: {csvnames}")
                else:
                    if "[Content_Types].xml" not in names:
                        raise ValueError(f"{key}: not an XLSX zip")
                    try:
                        import openpyxl
                    except ImportError as exc:
                        raise RuntimeError("Install openpyxl for Coates workbook introspection") from exc
                    wb=openpyxl.load_workbook(io.BytesIO(content),read_only=True,data_only=True)
                    fileinfo["worksheets"]={}
                    for ws in wb:
                        row=next(ws.values,None)
                        fileinfo["worksheets"][ws.title]={
                           "total_rows_including_header":ws.max_row,
                           "total_cols":ws.max_column,
                           "first_row_values": [str(x) if x is not None else "" for x in (row or [])][:90]
                        }
                    wb.close()
            receipt["original_files"][key]=fileinfo
            print("SOURCE_AUDIT",key, fileinfo["bytes"],fileinfo["sha256"])
            for name,m in fileinfo.get("members",{}).items():
                print("SOURCE_CSV",name,m["rows"],"columns",len(m["columns"]))
                if key=="crall_zenodo_zip":print("SOURCE_FIELDS",name,repr(m["columns"]))
            for name,m in fileinfo.get("worksheets",{}).items():
                print("SOURCE_SHEET",key,name,m["total_rows_including_header"],m["total_cols"])
        except Exception as exc:
            receipt["source_errors"][key]=str(exc)[:1000]
            print("SOURCE_UNAVAILABLE",key,str(exc)[:600])
    receipt["status"] = ("source_only_partial_or_full" if receipt["original_files"]
                          else "source_intake_blocked")
    Path(args.output).parent.mkdir(parents=True,exist_ok=True)
    Path(args.output).write_text(json.dumps(receipt,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    if not receipt["original_files"]:
        raise SystemExit("TNOA_WIND_ECOLOGY_SOURCE_INTAKE_BLOCKED: no raw dataset was accessible; no ecological results produced")
    print("TNOA_WIND_ECOLOGY_SOURCE_INTAKE_PASS: verified",len(receipt["original_files"]),
          "sources, unavailable",len(receipt["source_errors"]),
          "— no outcome statistics")

if __name__=="__main__":
    main()
